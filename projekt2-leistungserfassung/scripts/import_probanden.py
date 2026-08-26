"""Importiert die ECHTEN (pseudonymisierten) Ist-Daten der 5 Probanden.

    python -m scripts.import_probanden

Quelle: "Dokumente/AuswertungKursleistungenIKIV SoSe26_korrigiert.xlsx",
Blätter IK I–IK IV – DIREKT geparst (openpyxl, data_only).

Warum direkt aus der Excel (seit 2026-08-10, v2): Der Umweg über die
Prototyp-Datenbank (ki-kons-carolinum-2026, generate_test_sql.py) verlor
systematisch alle Leistungen, deren Spalte KEINE eigene Anzahl-Spalte hat
(Krone, Brückenglied, Teilprothese, Teleskop, MEG, Interimsprothese,
Unterfütterung, Proth. Recall, WF-Revision, Aufbissbehelf, Prophylaxe Kind,
Non-invasiv/invasiv Kind, Inlay/Teilkrone, Kronenrandfensterung, …) bzw.
verschob sie um eine Spalte – z. B. Stud. 1 IK I: Excel 71,5 P, alte DB 51,5 P.

Excel-Spalten-Layout (Zeile 4 = Leistungsname, Zeile 5 = Unterüberschrift):
- Paar-Klassen:   [Name+"Anzahl"] [ "Punkte" ]   (z. B. AIT, UPT, Füllungen)
- Einzel-Klassen: [Name+"Punkte"]                (z. B. Krone, WF-Revision)
- Block-Summen ("Gesamt") und Hospitationen (ohne Punkte-Label) werden
  übersprungen; Hosp. zählt laut Blatt nicht in die Semestersumme.

Was passiert:
- Datenbank wird NEU aufgebaut; danach stehen die 5 Probanden
  "Stud. 1"–"Stud. 5" mit ihren echten Kursleistungen drin.
- Ein Eintrag pro (Proband, Leistung, Semester); points = Semestersumme,
  count = 1 (App rechnet points*count), erfasste Anzahl steht in der Notiz.
- Validierung: Semestersumme der Einträge muss exakt der "IK N Gesamt"-Spalte
  des Blattes entsprechen (die Hosp. ausschließt). Zusätzlich Abgleich mit dem
  GESAMT-Blatt (dort rechnet UPT mit 2,5 P pauschal -> kleine bekannte
  Rundungsdifferenz möglich, wird ausgewiesen).

Demo-Daten zurückholen: python -m scripts.init_db
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from werkzeug.security import generate_password_hash

from app import create_app
from app.catalog import find_item
from app.db import get_db, init_db
from app.security import new_token

XLSX = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "Dokumente",
    "AuswertungKursleistungenIKIV SoSe26_korrigiert.xlsx",
)

# Excel-Leistungsname (Zeile 4) -> Katalog-Code der Blauen Liste
NAME_MAP = {
    "AIT": "pa_ait",
    "UPT": "pa_upt",
    "Klasse I und II": "zhs_fuell_1_2",
    "Klasse III und IV": "zhs_fuell_3_4",
    "Klasse V": "zhs_fuell_5",
    "Befund": "zhs_befund",
    "PZR": "zhs_prophylaxe",
    "WK": "endo_aufbereitung",
    "WF": "endo_fuellung",
    "WF-Revision": "endo_revision",
    "Befund Kind": "kzh_befund",
    "Prophylaxe Kind": "kzh_prophylaxe",
    "Non-invasiv/invasiv Kind": "kzh_behandlung",
    "Krone": "proth_krone",
    "Brückenglied": "proth_bruecke",
    "Totalprothese": "proth_totalprothese",
    "Teilprothese Doppelkrone": "proth_teilprothese",
    "Teleskop": "proth_teleskop",
    "MEG": "proth_modellguss",
    "Interimsprothese": "proth_interims",
    "Unterfütterung": "proth_unterfuett",
    "Remontage": "proth_remontage",
    "Proth. Recall": "proth_recall",
    "Aufbissbehelf": "proth_aufbiss",
    "Inlay/Teilkrone": "schnitt_inlay",
    "Kronenrandfensterung": "schnitt_reparatur",
    "Stumpfaufbau(KVB)": "schnitt_stumpfaufbau",
    "Stiftaufbau": "schnitt_stiftaufbau",
}

# Excel-Block -> (Gesamt-Spaltenlabel in Zeile 4, Leistungen des Blocks,
#                 Katalog-Code für nicht aufgeschlüsselte Block-Restwerte).
# Hintergrund: In Blatt IK III steht festsitzende Prothetik teils NUR als
# Blocksumme ("Festsitzend Gesamt"), ohne Einzel-Leistungen (Stud. 1: 15 P,
# Stud. 5: 99 P). Solche Restwerte werden als eigener, markierter Eintrag
# importiert, damit die Semestersumme stimmt.
BLOCKS = [
    ("Parodontologie", ["AIT", "UPT"], "pa_ait"),
    (
        "ZHS/Präv./Rest.",
        ["Klasse I und II", "Klasse III und IV", "Klasse V", "Befund", "PZR"],
        "zhs_fuell_1_2",
    ),
    ("Endodontologie", ["WK", "WF", "WF-Revision"], "endo_aufbereitung"),
    (
        "Kinderzahnheilkunde",
        ["Befund Kind", "Prophylaxe Kind", "Non-invasiv/invasiv Kind"],
        "kzh_behandlung",
    ),
    ("Festsitzend", ["Krone", "Brückenglied"], "proth_krone"),
    (
        "Herausnehmbar",
        [
            "Totalprothese",
            "Teilprothese Doppelkrone",
            "Teleskop",
            "MEG",
            "Interimsprothese",
            "Unterfütterung",
            "Remontage",
            "Proth. Recall",
        ],
        "proth_teilprothese",
    ),
    (
        "Schnittmenge",
        ["Inlay/Teilkrone", "Kronenrandfensterung", "Stumpfaufbau(KVB)", "Stiftaufbau"],
        "schnitt_stumpfaufbau",
    ),
]

SHEETS = ["IK I", "IK II", "IK III", "IK IV"]
SEM_DATE = {
    "IK I": "2024-11-15",
    "IK II": "2025-05-20",
    "IK III": "2025-11-18",
    "IK IV": "2026-05-12",
}


def detect_columns(ws):
    """Liefert [(excel_name, anzahl_col|None, punkte_col)] nach Layout-Regeln."""
    out = []
    for c in range(1, ws.max_column + 1):
        name = ws.cell(4, c).value
        name = name.strip() if isinstance(name, str) else None
        if name not in NAME_MAP:
            continue
        sub = ws.cell(5, c).value
        nxt = ws.cell(5, c + 1).value if c + 1 <= ws.max_column else None
        if sub == "Anzahl" and nxt == "Punkte":
            out.append((name, c, c + 1))
        elif sub == "Punkte":
            out.append((name, None, c))
    return out


def sheet_total_col(ws, sheet):
    """Spalte der 'IK N Gesamt'-Summe (Zeile 4 beginnt mit Blattnamen, Zeile 5 'Gesamt')."""
    for c in range(1, ws.max_column + 1):
        h4 = ws.cell(4, c).value
        if (
            isinstance(h4, str)
            and h4.strip().startswith(sheet)
            and ws.cell(5, c).value == "Gesamt"
        ):
            return c
    return None


def block_total_cols(ws):
    """Block-Label -> Spalte der Block-'Gesamt'-Zelle."""
    labels = {b[0] for b in BLOCKS}
    out = {}
    for c in range(1, ws.max_column + 1):
        h4 = ws.cell(4, c).value
        if (
            isinstance(h4, str)
            and h4.strip() in labels
            and ws.cell(5, c).value == "Gesamt"
        ):
            out[h4.strip()] = c
    return out


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


wb = openpyxl.load_workbook(XLSX, data_only=True)

app = create_app()
with app.app_context():
    init_db()
    db = get_db()

    cur = db.execute(
        "insert into users (username, password_hash, role, real_name) values (?,?,?,?)",
        (
            "doktorand1",
            generate_password_hash("doktorand1"),
            "doktorand",
            "Dr. Sabine Wagner",
        ),
    )
    dok_id = cur.lastrowid

    student_ids = {}
    print("== Probanden-Zugänge (Testbetrieb: Login = Passwort) ==")
    for n in range(1, 6):
        pseudonym = "S-" + new_token(4)
        cur = db.execute(
            "insert into students (pseudonym, real_name, cohort) values (?,?,?)",
            (pseudonym, f"Stud. {n}", "Probanden-Kohorte IK I–IV (Echtdaten)"),
        )
        student_ids[n] = cur.lastrowid
        # Testbetrieb: einfacher Login stud1..stud5 (Passwort = Benutzername).
        # Vor einem Echtbetrieb MUESSEN individuelle Passwoerter gesetzt werden.
        username = f"stud{n}"
        db.execute(
            "insert into users (username, password_hash, role, student_id) values (?,?,?,?)",
            (username, generate_password_hash(username), "student", student_ids[n]),
        )
        print(f"  Stud. {n}: Login {username} / {username}")

    imported = 0
    errors = []
    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        cols = detect_columns(ws)
        tcol = sheet_total_col(ws, sheet)
        btcols = block_total_cols(ws)

        def insert_row(sid, code, punkte, note, sheet=sheet):
            it = find_item(code)
            date = SEM_DATE[sheet]
            db.execute(
                "insert into performances (public_token, student_id, semester, "
                "category, item_code, item_name, count, points, patient_ref, "
                "performed_on, note, status, created_at, submitted_at, "
                "reviewed_by, reviewed_at, review_comment) "
                "values (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    new_token(),
                    student_ids[sid],
                    sheet,
                    it["category"],
                    it["code"],
                    it["name"],
                    1,
                    punkte,
                    None,
                    date,
                    note,
                    "approved",
                    date + "T09:00:00",
                    date + "T10:00:00",
                    dok_id,
                    date + "T15:00:00",
                    "Import aus AuswertungKursleistungenIKIV SoSe26_korrigiert.xlsx",
                ),
            )

        for row in range(6, 11):
            name = ws.cell(row, 1).value
            if not name or not str(name).replace(" ", "").startswith("Stud."):
                continue
            sid = int(str(name).split(".")[1])
            sem_sum = 0.0
            per_name = {}
            for excel_name, ca, cp in cols:
                punkte = num(ws.cell(row, cp).value)
                anzahl = num(ws.cell(row, ca).value) if ca else None
                per_name[excel_name] = punkte
                if punkte == 0 and not anzahl:
                    continue
                note = f"Import Kursauswertung Blatt {sheet} (Semestersumme)"
                if anzahl:
                    note += f", erfasste Anzahl: {anzahl:g}"
                insert_row(sid, NAME_MAP[excel_name], punkte, note)
                sem_sum += punkte
                imported += 1
            # Block-Restwerte: Blocksumme > Summe der Einzel-Leistungen
            # (kommt vor, wenn das Kurs-Team nur die Blocksumme eintrug).
            for label, members, rest_code in BLOCKS:
                if label not in btcols:
                    continue
                block_total = num(ws.cell(row, btcols[label]).value)
                item_sum = sum(per_name.get(m, 0.0) for m in members)
                rest = block_total - item_sum
                if rest > 1e-6:
                    insert_row(
                        sid,
                        rest_code,
                        rest,
                        f"Import Blatt {sheet}: Blocksumme '{label}' ohne "
                        f"Einzelaufschlüsselung ({rest:g} P Restwert)",
                    )
                    sem_sum += rest
                    imported += 1
                elif rest < -1e-6:
                    errors.append(
                        f"{sheet} Stud. {sid}: Block '{label}' Einzelwerte {item_sum:g} "
                        f"> Blocksumme {block_total:g}"
                    )
            # Validierung 1: Semestersumme == "IK N Gesamt"-Zelle des Blattes
            expected = num(ws.cell(row, tcol).value) if tcol else None
            if expected is not None and abs(sem_sum - expected) > 1e-6:
                errors.append(
                    f"{sheet} Stud. {sid}: importiert {sem_sum:g} != Blatt-Gesamt {expected:g}"
                )

    db.commit()
    print(f"\n{imported} Leistungszeilen importiert (Quelle: Excel-Blätter direkt).")

    print("\n== Validierung 1: Semestersummen vs. Blatt-Gesamt ==")
    if errors:
        for e in errors:
            print("  FEHLER:", e)
    else:
        print("  alle Semestersummen exakt gleich den Blatt-Gesamt-Zellen: OK")

    print("\n== Validierung 2: Gesamtpunkte je Proband vs. GESAMT-Blatt (AU) ==")
    gs = wb["GESAMT"]
    gcol = None
    for c in range(1, gs.max_column + 1):
        if (
            isinstance(gs.cell(4, c).value, str)
            and gs.cell(4, c).value.strip().startswith("GESAMT")
            and gs.cell(5, c).value == "Gesamt"
        ):
            gcol = c
    ok = True
    for row in range(6, 11):
        name = gs.cell(row, 1).value
        if not name:
            continue
        sid = int(str(name).split(".")[1])
        expected = num(gs.cell(row, gcol).value)
        got = db.execute(
            "select coalesce(sum(points*count),0) from performances where student_id=?",
            (student_ids[sid],),
        ).fetchone()[0]
        diff = got - expected
        mark = (
            "OK"
            if abs(diff) < 1e-6
            else f"Abweichung {diff:+g} (bekannte UPT-Rundung im GESAMT-Blatt)"
            if abs(diff) <= 2
            else "FEHLER"
        )
        if mark == "FEHLER":
            ok = False
        print(f"  Stud. {sid}: App={got:g}  GESAMT-Blatt={expected:g}  {mark}")
    if errors or not ok:
        print("\nVALIDIERUNG: FEHLGESCHLAGEN")
        sys.exit(1)
    print("\nVALIDIERUNG: BESTANDEN")
