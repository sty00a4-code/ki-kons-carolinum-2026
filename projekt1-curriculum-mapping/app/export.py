"""Export eines Ergebnisses als Excel-Datei.

Das Blatt "Zuordnung" enthält Prüfziel, Abdeckung, Begründung und Lernziel-ID
sowie Teil, Wortlaut der Lernziele und Hinweise. Die Datei lässt sich auch
nach manuellen Korrekturen wieder importieren.
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import formate, ki_dienst, zuordnung

ABDECKUNG = {zuordnung.ABGEDECKT: "X", zuordnung.TEILWEISE: "teilweise", zuordnung.NICHT: ""}

_KOPF_FUELLUNG = PatternFill("solid", fgColor="E6EAF7")
_GRUPPE_FUELLUNG = {
    "teil": PatternFill("solid", fgColor="D5DCF2"),
    "bereich": PatternFill("solid", fgColor="EEF1FB"),
    "gruppe": PatternFill("solid", fgColor="F7F8FD"),
}
_UMBRUCH = Alignment(wrap_text=True, vertical="top")


def dateiname(auswertung):
    titel = re.sub(r"[^\w.,() -]+", " ", auswertung["titel"]).strip()
    titel = re.sub(r"\s+", " ", titel)[:80]
    return f"Zuordnung {titel}.xlsx"


def xlsx(auswertung, ergebnis, liste):
    mappe = Workbook()
    _zuordnung(mappe.active, ergebnis)
    _lernziele(mappe.create_sheet("Nach Lernzielen"), ergebnis)
    _angaben(mappe.create_sheet("Angaben"), auswertung, ergebnis, liste)
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def _zuordnung(blatt, ergebnis):
    blatt.title = "Zuordnung"
    kopf = [
        "Teil",
        "Prüfziel",
        "Abdeckung",
        "Begründung",
        "Lernziel-ID (Kapitel)",
        "Wortlaut der Lernziele",
        "Hinweis",
    ]
    _kopfzeile(blatt, kopf, [7, 48, 11, 70, 22, 70, 30])
    for t in ergebnis.tabelle:
        art, e = t["art"], t["eintrag"]
        if art in _GRUPPE_FUELLUNG:
            name = e["name"] if art == "teil" else f"{e['nr']} {e['name']}"
            _zeile(blatt, [e.get("teil") or "", name], fuellung=_GRUPPE_FUELLUNG[art], fett=True)
        elif art == "fehlt":
            _zeile(
                blatt,
                [
                    e["teil"],
                    f"{e['nr']} {e['name']}",
                    "",
                    "",
                    "",
                    "",
                    "fehlt in der Antwort des Modells",
                ],
            )
        else:
            z = t["zeile"]
            wortlaut = "\n".join(
                f"{lid}: {ergebnis.lernziel_texte[lid]}"
                for lid in z.lernziele
                if lid in ergebnis.lernziel_texte
            )
            hinweise = [] if e else ["nicht in der Prüfziel-Liste"]
            unbekannt = [lid for lid in z.lernziele if lid in ergebnis.unbekannte_ids]
            if unbekannt:
                hinweise.append("nicht im Lernziel-Dokument: " + ", ".join(unbekannt))
            _zeile(
                blatt,
                [
                    e["teil"] if e else (z.teil or ""),
                    f"{e['nr']} {e['name']}" if e else f"{z.nr} {z.name}",
                    ABDECKUNG[z.stufe],
                    z.begruendung,
                    ", ".join(z.lernziele),
                    wortlaut,
                    "; ".join(hinweise),
                ],
            )
    blatt.auto_filter.ref = f"A1:{get_column_letter(len(kopf))}{blatt.max_row}"


def _lernziele(blatt, ergebnis):
    _kopfzeile(
        blatt, ["Lernziel-ID", "Wortlaut", "Anzahl Prüfziele", "Prüfziele"], [14, 70, 10, 70]
    )
    for lz in ergebnis.lernziele:
        pruefziele = "\n".join(_pruefziel_text(z) for z in lz["zeilen"])
        hinweis = " (kommt im Lernziel-Dokument nicht vor)" if lz["unbekannt"] else ""
        _zeile(blatt, [lz["id"], (lz["text"] or "") + hinweis, lz["anzahl"], pruefziele])


def _pruefziel_text(z):
    """Liefert Teil, Nummer und Name, bei zugeordneten Zeilen aus der Prüfziel-Liste."""
    if z.eintrag:
        return f"{z.eintrag['teil']} {z.eintrag['nr']} {z.eintrag['name']}"
    return f"{z.teil or ''} {z.nr} {z.name}".strip()


def _angaben(blatt, auswertung, ergebnis, liste):
    blatt.column_dimensions["A"].width = 26
    blatt.column_dimensions["B"].width = 100
    temperatur = auswertung["temperatur"]
    erstellt = formate.datum(auswertung["erstellt_am"])
    if auswertung["erstellt_von"]:
        erstellt += f" von {auswertung['erstellt_von']}"
    werte = [
        ("Titel", auswertung["titel"]),
        ("Kategorie", auswertung["kategorie"]),
        ("Prüfziel-Liste", liste.titel),
        ("Modell", ki_dienst.modell_name(auswertung["modell"])),
        ("Temperatur", "" if temperatur is None else temperatur),
        (
            "Quelle",
            "KI-Durchlauf in der App"
            if auswertung["quelle"] == "ki"
            else auswertung["quelle_hinweis"],
        ),
        ("Vorlesung", auswertung["vorlesung_name"] or ""),
        ("Lernziele", auswertung["lernziele_name"] or ""),
        ("Erstellt", erstellt),
        ("Zeilen in der Antwort", len(ergebnis.zeilen)),
        ("Abgedeckt (X)", ergebnis.abgedeckt),
        ("Teilweise abgedeckt", ergebnis.teilweise),
        ("Nicht abgedeckt", ergebnis.nicht),
        ("Prüfziele ohne Zeile", len(ergebnis.fehlend)),
        ("Prompt", auswertung["prompt"] or ""),
    ]
    for name, wert in werte:
        _zeile(blatt, [name, wert])
        blatt.cell(blatt.max_row, 1).font = Font(bold=True)


def _kopfzeile(blatt, spalten, breiten):
    blatt.append(spalten)
    for i, breite in enumerate(breiten, 1):
        zelle = blatt.cell(1, i)
        zelle.font = Font(bold=True)
        zelle.fill = _KOPF_FUELLUNG
        blatt.column_dimensions[get_column_letter(i)].width = breite
    blatt.freeze_panes = "A2"


def _zeile(blatt, werte, *, fuellung=None, fett=False):
    blatt.append([_text(w) for w in werte])
    for zelle in blatt[blatt.max_row]:
        if isinstance(zelle.value, str) and zelle.value.startswith("="):
            # Text aus der Antwort nie als Formel auswerten.
            zelle.data_type = "s"
        zelle.alignment = _UMBRUCH
        if fuellung:
            zelle.fill = fuellung
        if fett:
            zelle.font = Font(bold=True)


def _text(wert):
    if isinstance(wert, str):
        # Steuerzeichen sind in XLSX-Dateien nicht erlaubt.
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", wert)
    return wert
