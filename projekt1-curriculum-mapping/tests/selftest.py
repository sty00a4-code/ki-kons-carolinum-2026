"""Selbsttest der Anwendung, ohne Netzwerkzugriff.

    python -m tests.selftest

Abgedeckt sind Login und CSRF-Schutz, das Lesen von Modellantworten, ein
vollständiger Durchlauf gegen einen lokalen Test-Server für den KI-Dienst
(Upload, Streaming, Tabelle, Ansichten, Excel-Export und Reimport, Vergleich,
erneutes Ausführen, Löschen) sowie Fehlerfälle und Grenzwerte.

Optional prüft der Test eine echte Ergebnis-Excel mit drei Blättern:

    P1_TEST_EXCEL="ModellAuswertung Mistral.xlsx" \\
    P1_TEST_LERNZIELE=NKLZLernziele242248.pdf python -m tests.selftest
"""

import http.server
import io
import json
import os
import re
import sys
import tempfile
import threading
import time
from dataclasses import dataclass, field

import openpyxl
from werkzeug.security import generate_password_hash

from app import create_app, db, dokumente, excel_import, export, ki_dienst, pruefziele, zuordnung
from app.db import get_db

MB = 1024 * 1024
TEST_EXCEL = os.environ.get("P1_TEST_EXCEL", "")
TEST_LERNZIELE = os.environ.get("P1_TEST_LERNZIELE", "")


@dataclass
class Protokoll:
    pruefungen: int = 0
    fehler: list = field(default_factory=list)


PROTOKOLL = Protokoll()


def pruefe(bedingung, beschreibung):
    PROTOKOLL.pruefungen += 1
    if not bedingung:
        PROTOKOLL.fehler.append(beschreibung)
        print("  FEHLER:", beschreibung)


# Testdateien


def mini_pdf(zeilen):
    """Erzeugt eine minimale PDF-Datei mit Textebene (Helvetica, nur ASCII)."""

    def maskieren(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    inhalt = (
        "BT /F1 10 Tf 40 800 Td 13 TL " + " ".join(f"({maskieren(z)}) '" for z in zeilen) + " ET"
    )
    objekte = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        f"<< /Length {len(inhalt)} >>\nstream\n{inhalt}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    daten = b"%PDF-1.4\n"
    versaetze = []
    for nr, objekt in enumerate(objekte, 1):
        versaetze.append(len(daten))
        daten += f"{nr} 0 obj\n{objekt}\nendobj\n".encode("latin-1")
    xref = len(daten)
    daten += f"xref\n0 {len(objekte) + 1}\n0000000000 65535 f \n".encode()
    for v in versaetze:
        daten += f"{v:010d} 00000 n \n".encode()
    daten += (
        f"trailer\n<< /Size {len(objekte) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return daten


LERNZIELE_PDF = mini_pdf(
    [
        "25 Orale Medizin und systemische Aspekte",
        "25.1 Die Absolventin kann Auswirkungen von Allgemeinerkrankungen darstellen.",
        (
            "25.1.1.1 wichtige relevante Erkrankungen erkennen und in der Therapieplanung "
            "beruecksichtigen. 2"
        ),
        "25.1.1.2 psychische Auffaelligkeiten erkennen und somatische Ursachen abklaeren. 3a",
        (
            "25.2.1.5 Wechselwirkungen von Arzneimitteln mit der zahnaerztlichen Behandlung "
            "beschreiben. 2"
        ),
        "25.3.1.1 Infektionskrankheiten und ihre Bedeutung fuer die Behandlung erlaeutern. 3a",
    ]
)
VORLESUNG_PDF = mini_pdf(
    [
        "Querschnittsbereich Z2: Orale Medizin und systemische Aspekte",
        "Diabetes mellitus, kardio- und zerebrovaskulaere Erkrankungen, Schwangerschaft",
        "Parodontitis als Manifestation einer systemischen Erkrankung",
    ]
)


def modellantwort(liste):
    """Erzeugt eine Antwort im Format des Prompts mit typischen Abweichungen.

    Jedes dritte Prüfziel ist nicht abgedeckt, eine Lernziel-ID ist unbekannt
    und eine Zeile gehört zu keinem Prüfziel der Liste.
    """
    zeilen = [
        "<think>Ich gehe die Liste Zeile für Zeile durch.</think>",
        "Prüfziele;Abdeckung;Begründung;Lernziel-ID_Kapitel",
    ]
    teil = None
    for i, e in enumerate(liste.zaehlend):
        if e["teil"] != teil:
            teil = e["teil"]
            zeilen.append(
                f"**{teil} {'Erkrankungen' if teil == 'VI' else 'Übergeordnete Kompetenzen'}**"
            )
        if i % 3 == 0:
            zeilen.append(f"{e['nr']} {e['name']}; ;Keine direkte Abdeckung im zweiten Dokument.;")
        else:
            ids = "25.1.1.1, 25.2.1.5" if i % 2 else "25.3.1.1"
            if i == 1:
                ids += ", 25.9.9.9"
            zeilen.append(f"{e['nr']} {e['name']};X;Begründung zu {e['name']}.;{ids}")
    zeilen.append("99.9.9 Erfundenes Prüfziel;X;Steht nicht in der Liste.;25.1.1.1")
    return "\n".join(zeilen)


# Test-Server für den KI-Dienst


ANFRAGEN = []


class KiDienst(http.server.BaseHTTPRequestHandler):
    """Beantwortet /chat/completions je nach modus gestreamt oder mit Fehler."""

    modus = "ok"
    antwort = ""

    def log_message(self, *args):
        pass

    def do_POST(self):
        laenge = int(self.headers.get("Content-Length", 0))
        body = json.loads(self.rfile.read(laenge))
        ANFRAGEN.append(
            {"pfad": self.path, "auth": self.headers.get("Authorization"), "body": body}
        )
        if KiDienst.modus == "401":
            self.send_response(401)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(b'{"error": {"message": "Invalid API key"}}')
            return
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream")
        self.end_headers()
        text = KiDienst.antwort
        grund = "length" if KiDienst.modus == "length" else "stop"
        if KiDienst.modus == "length":
            text = text[: len(text) // 2]
        stuecke = [text[i : i + 400] for i in range(0, len(text), 400)]
        for nr, stueck in enumerate(stuecke):
            wahl = {"index": 0, "delta": {"content": stueck}, "finish_reason": None}
            if nr == len(stuecke) - 1:
                wahl["finish_reason"] = grund
            self.wfile.write(f"data: {json.dumps({'choices': [wahl]})}\n\n".encode())
            self.wfile.flush()
            time.sleep(0.01)
        self.wfile.write(b"data: [DONE]\n\n")


def dienst_starten():
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), KiDienst)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


# Hilfen für den Test-Client


def csrf(client, pfad):
    html = client.get(pfad).get_data(as_text=True)
    treffer = re.search(r'name="_csrf_token" value="([^"]+)"', html)
    return treffer.group(1) if treffer else ""


def warten_bis_fertig(client, token, sekunden=20):
    ende = time.monotonic() + sekunden
    while time.monotonic() < ende:
        stand = client.get(f"/ergebnisse/{token}/stand").get_json()
        if stand["status"] != "laeuft":
            return stand
        time.sleep(0.1)
    return client.get(f"/ergebnisse/{token}/stand").get_json()


def token_aus(antwort):
    treffer = re.search(r"/ergebnisse/([A-Za-z0-9_-]{16})$", antwort.headers.get("Location", ""))
    return treffer.group(1) if treffer else None


def neue_zuordnung(client, **ueberschreiben):
    daten = {
        "_csrf_token": csrf(client, "/neu"),
        "prompt": "Liste mir alle Prüfziele auf und ordne die Lernziele zu.",
        "modell": ki_dienst.modelle()[0][0],
        "temperatur": "0,5",
        "liste": pruefziele.STANDARD,
        "titel": "",
        "vorlesung": (io.BytesIO(VORLESUNG_PDF), "Vorlesung_1.pdf"),
        "lernziele": (io.BytesIO(LERNZIELE_PDF), "Lernziele_Kapitel_25.pdf"),
    }
    daten.update(ueberschreiben)
    return client.post("/neu", data=daten, content_type="multipart/form-data")


def anzahl_auswertungen(app):
    with app.app_context():
        return get_db().execute("select count(*) from auswertungen").fetchone()[0]


# Tests


def test_leser(liste):
    print("Leser der Modellantwort")
    text = modellantwort(liste)
    ergebnis = zuordnung.auswerten(
        text,
        liste,
        "25.1.1.1 wichtige Erkrankungen erkennen. 2\n"
        "25.2.1.5 Wechselwirkungen. 2\n"
        "25.3.1.1 Infektionen. 3a",
    )
    pruefe(
        len(ergebnis.zeilen) == len(liste.zaehlend) + 1,
        f"alle Zeilen gelesen ({len(ergebnis.zeilen)})",
    )
    pruefe(not ergebnis.fehlend, "kein Prüfziel fehlt")
    pruefe(len(ergebnis.ohne_liste) == 1, "erfundenes Prüfziel steht außerhalb der Liste")
    pruefe(
        ergebnis.unbekannte_ids == {"25.9.9.9"},
        f"unbekannte ID erkannt ({ergebnis.unbekannte_ids})",
    )
    pruefe(not ergebnis.uebrig, f"keine unlesbaren Zeilen ({ergebnis.uebrig[:2]})")
    erwartet_nicht = sum(1 for i, _ in enumerate(liste.zaehlend) if i % 3 == 0)
    pruefe(
        ergebnis.nicht == erwartet_nicht,
        f"nicht abgedeckt gezählt ({ergebnis.nicht} statt {erwartet_nicht})",
    )

    varianten = """
| Prüfziel | Abdeckung | Begründung | Lernziel-ID |
|---|---|---|---|
| 1.1.1 Karies | X | Wird unter Allgemeinerkrankungen behandelt. | 25.1.1.1 |
- **1.1.2 Erosion, Abrasion, Attrition**; teilweise; nur am Rand; 25.2.1.5
VIII Übergeordnete Kompetenzen
VIII.4.5.7 Behinderung und Gesundheit;;*(Keine direkte Abdeckung)*
"""
    ergebnis = zuordnung.auswerten(varianten, liste)
    stufen = {z.eintrag["schluessel"]: z.stufe for z in ergebnis.zeilen if z.eintrag}
    pruefe(stufen.get("VI-1.1.1") == zuordnung.ABGEDECKT, "Tabelle mit | gelesen")
    pruefe(
        stufen.get("VI-1.1.2") == zuordnung.TEILWEISE,
        "Aufzählung mit Fettdruck und teilweise gelesen",
    )
    pruefe(
        stufen.get("VIII-4.5.7") == zuordnung.NICHT,
        "VIII. vor der Nummer und Hinweis in Klammern gelesen",
    )
    hinweis = next((z.begruendung for z in ergebnis.zeilen if z.nr == "4.5.7"), "")
    pruefe(
        hinweis == "Keine direkte Abdeckung", f"Hinweis ohne Sternchen und Klammern ({hinweis!r})"
    )

    abweichend = zuordnung.auswerten("4.12.2 Depression;X;passt;25.1.1.2", liste)
    pruefe(
        abweichend.zeilen
        and abweichend.zeilen[0].eintrag
        and abweichend.zeilen[0].eintrag["nr"] == "4.13.2",
        "Nummer aus der Prüfziel-Tabelle (4.12.2 Depression) findet 4.13.2",
    )
    pruefe(zuordnung.kappa([2, 0, 2, 0], [2, 0, 2, 0]) == 1.0, "Kappa 1 bei voller Übereinstimmung")
    k = zuordnung.kappa([2, 2, 0, 0], [2, 0, 2, 0])
    pruefe(k is not None and abs(k) < 1e-9, f"Kappa 0 bei Zufallsniveau ({k})")


def test_zugriff(client):
    print("Login und CSRF")
    for pfad in [
        "/",
        "/neu",
        "/import",
        "/ergebnisse",
        "/ergebnisse/abc",
        "/vergleich",
        "/pruefziele",
    ]:
        antwort = client.get(pfad)
        pruefe(
            antwort.status_code == 302 and "/login" in antwort.headers["Location"],
            f"{pfad} ohne Login leitet zur Anmeldung",
        )
    pruefe(
        client.post("/login", data={"username": "doktorand1", "password": "x"}).status_code == 400,
        "Login ohne CSRF-Token wird abgelehnt",
    )
    antwort = client.post(
        "/login",
        data={
            "username": "doktorand1",
            "password": "falsch",
            "_csrf_token": csrf(client, "/login"),
        },
    )
    pruefe(
        "Login fehlgeschlagen" in antwort.get_data(as_text=True), "falsches Passwort meldet Fehler"
    )
    antwort = client.post(
        "/login",
        data={
            "username": "doktorand1",
            "password": "richtig-und-lang",
            "_csrf_token": csrf(client, "/login"),
        },
    )
    pruefe(antwort.status_code == 302, "Login mit richtigem Passwort")
    pruefe(
        client.post("/neu", data={"prompt": "x"}).status_code == 400,
        "Formular ohne CSRF-Token wird abgelehnt",
    )
    for pfad in ["/neu", "/import", "/ergebnisse", "/vergleich", "/pruefziele"]:
        pruefe(client.get(pfad).status_code == 200, f"{pfad} lädt nach Login")
    pruefe(
        client.get("/ergebnisse/gibt-es-nicht").status_code == 404,
        "unbekanntes Ergebnis ergibt 404",
    )


def test_ohne_schluessel(client):
    print("Ohne Schlüssel")
    os.environ.pop("P1_LLM_API_KEY", None)
    html = client.get("/neu").get_data(as_text=True)
    pruefe("noch kein Schlüssel" in html, "Hinweis auf fehlenden Schlüssel")
    antwort = neue_zuordnung(client)
    pruefe(
        antwort.status_code == 200 and "kein Schlüssel" in antwort.get_data(as_text=True),
        "ohne Schlüssel startet nichts",
    )


def test_durchlauf(client, liste):
    print("Durchlauf gegen den Test-Server")
    KiDienst.antwort = modellantwort(liste)
    KiDienst.modus = "ok"

    antwort = neue_zuordnung(client)
    token = token_aus(antwort)
    pruefe(token is not None, f"Upload leitet zum Ergebnis ({antwort.status_code})")
    if not token:
        return None
    stand = warten_bis_fertig(client, token)
    pruefe(stand["status"] == "fertig", f"Durchlauf fertig ({stand})")
    anfrage = ANFRAGEN[-1]
    pruefe(anfrage["pfad"] == "/v1/chat/completions", "richtiger Endpunkt")
    pruefe(anfrage["auth"] == "Bearer test-schluessel", "Schlüssel im Header")
    body = anfrage["body"]
    pruefe(body["stream"] is True and body["temperature"] == 0.5, "Stream und Temperatur 0,5")
    inhalt = body["messages"][1]["content"]
    pruefe(
        inhalt.index("Dokument 1: Prüfziele")
        < inhalt.index("Dokument 2: Lernziele")
        < inhalt.index("Dokument 3: Vorlesung"),
        "Reihenfolge der Dokumente in der Anfrage",
    )
    pruefe(
        "25.1.1.1 wichtige relevante Erkrankungen" in inhalt, "Text der Lernziel-PDF in der Anfrage"
    )
    pruefe("Parodontitis als Manifestation" in inhalt, "Text der Vorlesungs-PDF in der Anfrage")
    pruefe(inhalt.rstrip().endswith("ordne die Lernziele zu."), "Prompt am Ende der Anfrage")

    html = client.get(f"/ergebnisse/{token}").get_data(as_text=True)
    zaehlend = len(liste.zaehlend)
    abgedeckt = sum(1 for i, _ in enumerate(liste.zaehlend) if i % 3) + 1
    pruefe(f'<div class="mv">{abgedeckt}</div>' in html, f"Kachel abgedeckt zeigt {abgedeckt}")
    pruefe(html.count('data-stufe="2"') == abgedeckt, "eine Tabellenzeile je abgedecktem Prüfziel")
    pruefe("Nicht in der Prüfziel-Liste" in html, "erfundenes Prüfziel unten in der Tabelle")
    pruefe("25.9.9.9" in html and "chip unbekannt" in html, "unbekannte ID markiert")
    pruefe("wichtige relevante Erkrankungen erkennen" in html, "Wortlaut des Lernziels aus der PDF")
    pruefe("Ich gehe die Liste" not in html, "Denk-Abschnitt nicht in der Tabelle")
    for ansicht in ["lernziele", "antwort", "eingaben"]:
        antwort = client.get(f"/ergebnisse/{token}?ansicht={ansicht}")
        pruefe(antwort.status_code == 200, f"Ansicht {ansicht} lädt")
    pruefe(
        "Vorlesung_1.pdf"
        in client.get(f"/ergebnisse/{token}?ansicht=eingaben").get_data(as_text=True),
        "Dateiname unter Eingaben",
    )
    pruefe(
        f"{zaehlend} Prüfziele" in client.get("/neu").get_data(as_text=True),
        "Formular nennt die Zahl der Prüfziele",
    )
    return token


def test_export(app, client, liste, token):
    print("Excel-Export und Reimport")
    antwort = client.get(f"/ergebnisse/{token}/excel")
    pruefe(
        antwort.status_code == 200
        and "attachment" in antwort.headers.get("Content-Disposition", ""),
        "Excel-Export",
    )
    mappe = openpyxl.load_workbook(io.BytesIO(antwort.data))
    pruefe(
        mappe.sheetnames == ["Zuordnung", "Nach Lernzielen", "Angaben"],
        f"Blätter im Export ({mappe.sheetnames})",
    )
    kopf = [z.value for z in mappe["Zuordnung"][1]]
    pruefe(
        kopf[:5] == ["Teil", "Prüfziel", "Abdeckung", "Begründung", "Lernziel-ID (Kapitel)"],
        f"Spalten im Export ({kopf})",
    )
    antwort = client.post(
        "/import",
        data={
            "_csrf_token": csrf(client, "/import"),
            "excel": (io.BytesIO(antwort.data), "Export.xlsx"),
            "lernziele": (io.BytesIO(LERNZIELE_PDF), "Lernziele_Kapitel_25.pdf"),
            "modell": ki_dienst.modelle()[0][0],
            "liste": pruefziele.STANDARD,
            "titel": "Rückimport",
            "prompt": "",
        },
        content_type="multipart/form-data",
    )
    rueck = token_aus(antwort)
    pruefe(
        rueck is not None, f"Rückimport eines Blatts leitet zum Ergebnis ({antwort.status_code})"
    )
    if not rueck:
        return
    with app.app_context():
        con = get_db()
        a1 = con.execute("select * from auswertungen where token = ?", (token,)).fetchone()
        a2 = con.execute("select * from auswertungen where token = ?", (rueck,)).fetchone()
        e1 = zuordnung.auswerten(a1["rohantwort"], liste, a1["lernziele_text"])
        e2 = zuordnung.auswerten(a2["rohantwort"], liste, a2["lernziele_text"])

    def kurz(e):
        return [
            ((z.eintrag or {}).get("schluessel"), z.stufe, z.lernziele, z.begruendung)
            for z in e.zeilen
        ]

    pruefe(kurz(e1) == kurz(e2), "Rückimport ergibt dieselbe Zuordnung")
    pruefe(a2["titel"] == "Rückimport · Zuordnung", f"Titel des Rückimports ({a2['titel']})")
    antwort = client.get(f"/vergleich?t={token}&t={rueck}")
    pruefe(
        antwort.status_code == 200 and "<strong>1,00</strong>" in antwort.get_data(as_text=True),
        "Vergleich mit sich selbst ergibt Kappa 1",
    )


def test_wiederholen(app, client, token):
    print("Erneut ausführen")
    antwort = client.post(
        f"/ergebnisse/{token}/wiederholen",
        data={
            "_csrf_token": csrf(client, "/neu"),
            "modell": ki_dienst.modelle()[1][0],
            "temperatur": "1",
            "prompt": "Neuer Prompt für den zweiten Lauf.",
        },
    )
    neu = token_aus(antwort)
    pruefe(neu is not None and neu != token, "Erneut ausführen legt ein neues Ergebnis an")
    if neu:
        pruefe(warten_bis_fertig(client, neu)["status"] == "fertig", "zweiter Durchlauf fertig")
        body = ANFRAGEN[-1]["body"]
        pruefe(
            body["model"] == ki_dienst.modelle()[1][0] and body["temperature"] == 1.0,
            "zweiter Lauf mit neuem Modell und Temperatur",
        )
        pruefe(
            "Parodontitis als Manifestation" in body["messages"][1]["content"],
            "zweiter Lauf mit gespeicherter Vorlesung",
        )

    vorher = anzahl_auswertungen(app)
    antwort = client.post(
        f"/ergebnisse/{token}/wiederholen",
        data={
            "_csrf_token": csrf(client, "/neu"),
            "modell": ki_dienst.modelle()[1][0],
            "temperatur": "5",
            "prompt": "Geänderter Prompt, der nicht verloren gehen darf.",
        },
    )
    html = antwort.get_data(as_text=True)
    pruefe(
        antwort.status_code == 200
        and "zwischen 0 und 2" in html
        and "Geänderter Prompt, der nicht verloren gehen darf." in html
        and 'value="5"' in html
        and f'<option value="{ki_dienst.modelle()[1][0]}" selected>' in html
        and anzahl_auswertungen(app) == vorher,
        "Fehler beim erneuten Ausführen behält die Eingaben",
    )


def test_entwurf(client, token):
    print("Entwurf im Browser")
    for pfad in ("/neu", "/import", f"/ergebnisse/{token}?ansicht=eingaben"):
        html = client.get(pfad).get_data(as_text=True)
        pruefe(
            all(
                teil in html
                for teil in (
                    "entwurf.js",
                    "entwurfEinrichten(",
                    "entwurf-hinweis",
                    "entwurf-verwerfen",
                )
            ),
            f"Entwurf im Browser auf {pfad}",
        )
    skript = client.get("/static/entwurf.js")
    pruefe(
        skript.status_code == 200 and b"localStorage" in skript.data,
        "entwurf.js wird ausgeliefert",
    )
    skript.close()


def test_dienstfehler(client):
    print("Fehler des KI-Dienstes")
    KiDienst.modus = "401"
    fehler_token = token_aus(neue_zuordnung(client))
    stand = warten_bis_fertig(client, fehler_token) if fehler_token else {}
    pruefe(
        stand.get("status") == "fehler" and "abgelehnt" in (stand.get("fehler") or ""),
        f"abgelehnter Schlüssel ({stand})",
    )
    pruefe(
        "fehlgeschlagen" in client.get(f"/ergebnisse/{fehler_token}").get_data(as_text=True),
        "Fehlerkarte auf der Ergebnisseite",
    )

    KiDienst.modus = "length"
    kurz_token = token_aus(neue_zuordnung(client))
    stand = warten_bis_fertig(client, kurz_token) if kurz_token else {}
    html = client.get(f"/ergebnisse/{kurz_token}").get_data(as_text=True)
    pruefe(
        stand.get("status") == "fertig" and "Längengrenze" in html,
        "abgeschnittene Antwort wird gemeldet",
    )
    pruefe("Ohne Zeile in der Antwort:" in html, "fehlende Prüfziele werden aufgezählt")
    KiDienst.modus = "ok"
    return fehler_token


def test_eingabefehler(client):
    print("Eingabefehler")
    antwort = neue_zuordnung(client, lernziele=(io.BytesIO(b"kein pdf"), "Lernziele.docx"))
    pruefe(
        "Bitte eine PDF- oder Textdatei" in antwort.get_data(as_text=True),
        "Datei ohne PDF wird abgelehnt",
    )
    antwort = neue_zuordnung(client, vorlesung=(io.BytesIO(b""), ""))
    pruefe(
        "Vorlesung: Bitte eine Datei auswählen" in antwort.get_data(as_text=True),
        "fehlende Vorlesung wird gemeldet",
    )
    antwort = neue_zuordnung(client, temperatur="3")
    pruefe("zwischen 0 und 2" in antwort.get_data(as_text=True), "Temperatur außerhalb 0 bis 2")
    antwort = neue_zuordnung(
        client, prompt="Mein eigener Prompt mit genug Text", modell="gibt-es-nicht"
    )
    html = antwort.get_data(as_text=True)
    pruefe(
        "Modell aus der Liste" in html and "Mein eigener Prompt" in html,
        "Fehler behält den eingegebenen Prompt",
    )


def test_gleichzeitig(app, client):
    print("Gleichzeitige und verwaiste Durchläufe")
    with app.app_context():
        con = get_db()
        jetzt, alt = db.jetzt(), db.vor_minuten(60)
        for nr, zeit in enumerate([jetzt, jetzt, alt]):
            con.execute(
                "insert into auswertungen (token, titel, modell, pruefziele_liste, quelle, "
                "status, erstellt_am, aktualisiert_am) "
                "values (?, 'Test', 'm', ?, 'ki', 'laeuft', ?, ?)",
                (f"laeuft-{nr}", pruefziele.STANDARD, zeit, zeit),
            )
        con.commit()
    antwort = neue_zuordnung(client)
    pruefe(
        "Es laufen schon" in antwort.get_data(as_text=True),
        "höchstens zwei Durchläufe gleichzeitig",
    )
    stand = client.get("/ergebnisse/laeuft-2/stand").get_json()
    pruefe(
        stand["status"] == "fehler" and "Abgebrochen" in stand["fehler"],
        "verwaister Durchlauf wird beendet",
    )
    html = client.get("/ergebnisse/laeuft-0").get_data(as_text=True)
    pruefe('id="laufend"' in html, "laufender Durchlauf zeigt den Fortschritt")
    with app.app_context():
        get_db().execute("delete from auswertungen where token like 'laeuft-%'")
        get_db().commit()


def test_loeschen(client, token):
    print("Löschen")
    antwort = client.post(
        f"/ergebnisse/{token}/loeschen", data={"_csrf_token": csrf(client, "/neu")}
    )
    pruefe(
        antwort.status_code == 302 and client.get(f"/ergebnisse/{token}").status_code == 404,
        "Löschen entfernt das Ergebnis",
    )


def test_grenzen(app, client):
    print("Grenzen")
    grenze = app.config["MAX_CONTENT_LENGTH"]
    app.config["MAX_CONTENT_LENGTH"] = 2 * MB
    try:
        antwort = neue_zuordnung(
            client, vorlesung=(io.BytesIO(b"%PDF-" + b"0" * (3 * MB)), "gross.pdf")
        )
    finally:
        app.config["MAX_CONTENT_LENGTH"] = grenze
    pruefe(
        antwort.status_code == 413 and "größer als 2 MB" in antwort.get_data(as_text=True),
        f"zu große Datei ({antwort.status_code})",
    )

    liste = pruefziele.liste()
    ergebnis = zuordnung.auswerten(
        '1.1.1 Karies;X;=HYPERLINK("http://example.org");25.1.1.1', liste
    )
    a = {
        "titel": "Formel",
        "kategorie": 1,
        "modell": "m",
        "temperatur": None,
        "quelle": "import",
        "quelle_hinweis": "x",
        "vorlesung_name": None,
        "lernziele_name": None,
        "erstellt_am": "2026-09-14T12:00:00",
        "erstellt_von": None,
        "prompt": None,
    }
    blatt = openpyxl.load_workbook(io.BytesIO(export.xlsx(a, ergebnis, liste)))["Zuordnung"]
    zellen = [
        z
        for reihe in blatt.iter_rows()
        for z in reihe
        if isinstance(z.value, str) and z.value.startswith("=")
    ]
    pruefe(
        zellen and all(z.data_type == "s" for z in zellen),
        "Formeln aus der Antwort bleiben Text im Export",
    )


def test_echte_dateien(liste):
    if not os.path.exists(TEST_EXCEL):
        print("Echte Dateien: übersprungen (P1_TEST_EXCEL nicht gesetzt)")
        return
    print("Echte Dateien")
    with open(TEST_EXCEL, "rb") as fh:
        blaetter = excel_import.blaetter(fh.read())
    lernziele = ""
    if os.path.exists(TEST_LERNZIELE):
        with open(TEST_LERNZIELE, "rb") as fh:
            lernziele = dokumente.aufraeumen(dokumente.pdf_text(fh.read()))
    ergebnisse = {name: zuordnung.auswerten(text, liste, lernziele) for name, text, _ in blaetter}
    # Kontrollzahlen der drei Blätter (Temperatur 0.0, 0.5 und 1.0)
    zeilen = [len(e.zeilen) for e in ergebnisse.values()]
    pruefe(zeilen == [143, 143, 141], f"Zeilen je Blatt ({zeilen})")
    pruefe(
        all(not e.uebrig and not e.ohne_liste for e in ergebnisse.values()),
        "alle Zeilen einem Prüfziel zugeordnet",
    )
    abgedeckt = [e.abgedeckt for e in ergebnisse.values()]
    pruefe(abgedeckt == [116, 124, 86], f"X je Blatt wie in der Excel ({abgedeckt})")
    if lernziele:
        unbekannt = [sorted(e.unbekannte_ids) for e in ergebnisse.values()]
        pruefe(
            unbekannt[2] == ["4.5.2", "4.5.2.4", "4.5.2.5", "4.5.2.6"],
            f"unbekannte IDs im Blatt 1.0 ({unbekannt[2]})",
        )


def ablauf(app):
    """Führt alle Prüfungen gegen eine frisch angelegte Anwendung aus."""
    with app.app_context():
        get_db().execute(
            "insert into users (username, password_hash, role) values (?, ?, 'doktorand')",
            ("doktorand1", generate_password_hash("richtig-und-lang")),
        )
        get_db().commit()
    client = app.test_client()
    liste = pruefziele.liste()
    for var in ("P1_LLM_URL", "P1_LLM_API_KEY", "P1_LLM_MODELLE", "P1_LLM_MAX_TOKENS"):
        os.environ.pop(var, None)

    test_leser(liste)
    test_zugriff(client)
    test_ohne_schluessel(client)

    server = dienst_starten()
    os.environ["P1_LLM_URL"] = f"http://127.0.0.1:{server.server_address[1]}/v1"
    os.environ["P1_LLM_API_KEY"] = "test-schluessel"
    ki_dienst.ZWISCHENSTAND_SEKUNDEN = 0
    token = test_durchlauf(client, liste)
    if token:
        test_export(app, client, liste, token)
        test_wiederholen(app, client, token)
        test_entwurf(client, token)
    fehler_token = test_dienstfehler(client)
    test_eingabefehler(client)
    test_gleichzeitig(app, client)
    test_loeschen(client, fehler_token)
    server.shutdown()

    test_grenzen(app, client)
    test_echte_dateien(liste)


def main():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        ablauf(create_app({"DATABASE": os.path.join(tmp, "selftest.db"), "TESTING": True}))
    print()
    if PROTOKOLL.fehler:
        print(f"{len(PROTOKOLL.fehler)} von {PROTOKOLL.pruefungen} Prüfungen fehlgeschlagen.")
        sys.exit(1)
    print(f"Alle {PROTOKOLL.pruefungen} Prüfungen bestanden.")


if __name__ == "__main__":
    main()
